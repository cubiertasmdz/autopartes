#!/usr/bin/env python3
"""Herramienta del catálogo de paragolpes, parrillas y ópticas usadas.

Uso:  .venv/bin/python tools/catalogo.py <comando> [opciones]

Comandos:
  ver [texto] [--todos]            Lista productos (con datos internos). --todos incluye ocultos.
  nuevo <archivo.json>             Carga un producto nuevo (ingreso). Imprime el código asignado.
  sumar <id> <n> [--costo X]       Suma unidades a un lote existente (ingreso).
  fotos <id> <img...> [--reemplazar]  Comprime y agrega fotos al producto.
  venta <id> <n> [--monto X] [--nota "..."]
  precio <id> [--unidad X] [--par X|no]   (par = izquierda + derecha)
  baja <id> [--nota "..."]         Oculta el producto del catálogo.
  alta <id>                        Vuelve a mostrar un producto oculto.
  editar <id> campo=valor ...      Corrige datos (valor en JSON o texto).
  build                            Genera los archivos públicos (data/, p/).
  resumen [--mes AAAA-MM]
  planilla                         Regenera PLANILLA STOCK.xlsx (se hace sola en cada movimiento).
  excel                            Copia fechada de la planilla en exportes/.
  publicar "mensaje"               build + commit + push.
  limpiar-pruebas                  Borra los productos marcados como prueba.
"""
import datetime as dt
import html
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PRIV = ROOT / "privado"
STOCK = PRIV / "stock.json"
HIST = PRIV / "historial.json"
CONFIG = json.loads((ROOT / "tools" / "config.json").read_text())
FOTOS = ROOT / "fotos"
FONTS = ROOT / "tools" / "fonts"

# Campos que NUNCA salen al sitio público.
PRIVADOS = {"costo", "notas_privadas", "proveedor", "prueba"}


# ---------- datos ----------

def hoy():
    return dt.date.today()


def ahora():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def cargar():
    if not STOCK.exists():
        return {"siguiente": {}, "productos": []}
    return json.loads(STOCK.read_text())


def guardar(stock):
    PRIV.mkdir(exist_ok=True)
    STOCK.write_text(json.dumps(stock, ensure_ascii=False, indent=2) + "\n")


def historial():
    return json.loads(HIST.read_text()) if HIST.exists() else []


def registrar(tipo, prod, cantidad=None, monto=None, detalle=""):
    h = historial()
    h.append({
        "fecha": ahora(), "tipo": tipo, "id": prod["id"], "producto": nombre_interno(prod),
        "cantidad": cantidad, "monto": monto, "detalle": detalle,
        "costo_unit": prod.get("costo"), "precio_unit": prod.get("precio_unidad"),
        **({"prueba": True} if prod.get("prueba") else {}),
    })
    PRIV.mkdir(exist_ok=True)
    HIST.write_text(json.dumps(h, ensure_ascii=False, indent=2) + "\n")
    planilla()


def categoria(cat_id):
    for c in CONFIG["categorias"]:
        if c["id"] == cat_id:
            return c
    sys.exit(f"Categoría desconocida: {cat_id}")


def buscar(stock, pid):
    for p in stock["productos"]:
        if p["id"].upper() == pid.upper():
            return p
    sys.exit(f"No existe el producto {pid}")


def pesos(n):
    return "$" + f"{int(round(n)):,}".replace(",", ".") if n is not None else "—"


# ---------- textos del producto ----------

def anios(p):
    """'2008 a 2012' / 'desde 2015' / '2010' según anio_desde / anio_hasta."""
    d, h = p.get("anio_desde"), p.get("anio_hasta")
    if d and h and d != h:
        return f"{d} a {h}"
    if d and h is None and p.get("anio_en_adelante"):
        return f"desde {d}"
    return str(d or h or "")


def lado_txt(p):
    """Paragolpes: 'delantero'/'trasero'. Ópticas: 'delantera izquierda', etc. Parrillas: nada."""
    pos = (p.get("posicion") or "").lower()
    lado = (p.get("lado") or "").lower()
    if p["tipo"] == "paragolpes":
        return pos
    if p["tipo"] == "optica":
        if lado == "par":  # izquierda + derecha juntas
            return {"delantero": "delanteras", "trasero": "traseras"}.get(pos, pos)
        pos = {"delantero": "delantera", "trasero": "trasera"}.get(pos, pos)
        lado = {"izquierdo": "izquierda", "derecho": "derecha"}.get(lado, lado)
        return " ".join(x for x in [pos, lado] if x)
    return ""


def vehiculo(p):
    return " ".join(str(x) for x in [p.get("marca"), p.get("modelo")] if x)


def nombre_pieza(p):
    """'Paragolpes delantero', 'Óptica trasera derecha', 'Parrilla'."""
    cat = categoria(p["tipo"])
    if es_par(p):
        return " ".join(x for x in ["Par de", cat["plural"], lado_txt(p)] if x)
    return " ".join(x for x in [cat["singular"].capitalize(), lado_txt(p)] if x)


def es_par(p):
    return p["tipo"] == "optica" and p.get("lado") == "par"


def nombre_interno(p):
    return " ".join(x for x in [nombre_pieza(p), vehiculo(p), anios(p)] if x)


def nombre_planilla(p, n):
    if p.get("nombre"):
        return p["nombre"]
    return f"{n} x {nombre_interno(p)}"


def visible(p):
    return p.get("activo", True) and int(p.get("cantidad", 0)) > 0


def etiqueta_cantidad(p, n):
    if es_par(p):
        return "Par (izq. y der.)"
    if n == 1:
        return "1 unidad"
    return f"{n} unidades"


def precios_extra(p, n):
    """Precio por el par (izquierda + derecha), solo si hay 2 o más."""
    if p.get("precio_par") and n >= 2:
        return [{"txt": "el par", "monto": p["precio_par"]}]
    return []


def numeros(p):
    """Años que cubre la pieza: buscar 'gol 2010' encuentra un paragolpes 2008 a 2012."""
    d, h = p.get("anio_desde"), p.get("anio_hasta")
    if d and not h and p.get("anio_en_adelante"):
        h = hoy().year
    d, h = d or h, h or d
    nums = set(str(a) for a in range(int(d), int(h) + 1)) if d else set()
    for c in p.get("claves", []):
        nums.update(re.findall(r"\d+", str(c)))
    if p.get("numero_pieza"):
        nums.add(re.sub(r"\W", "", str(p["numero_pieza"])).lower())
    return sorted(nums)


# ---------- versión pública de cada producto ----------

def publico(p):
    cat = categoria(p["tipo"])
    n = int(p["cantidad"])
    etiquetas = []
    titulo = p.get("titulo_web") or " ".join(x for x in [nombre_pieza(p), vehiculo(p)] if x)
    partes = []
    if anios(p):
        partes.append(anios(p).capitalize())
    if p.get("version"):
        partes.append(p["version"])
    sub = p.get("subtitulo") or " · ".join(partes)

    if p.get("original") is True:
        etiquetas.append("Original")
    elif p.get("original") is False:
        etiquetas.append("Alternativa")
    if p["tipo"] == "optica" and p.get("tecnologia"):
        etiquetas.append(p["tecnologia"])  # "Halógena", "LED", "Xenón", "Con lupa"...
    if p.get("color"):
        etiquetas.append(p["color"])
    etiquetas += p.get("detalles_estado", [])
    etiquetas += p.get("etiquetas", [])
    if p.get("numero_pieza"):
        etiquetas.append(f'N.º de pieza {p["numero_pieza"]}')

    sirve = ", ".join(p.get("vehiculos", []))
    pieza = " ".join(x for x in [cat["singular"], lado_txt(p)] if x)
    if es_par(p):
        cuanto = "el " + nombre_pieza(p).lower()
    elif n == 1:
        cuanto = f'{cat["articulo"]} {pieza}'
    else:
        cuanto = f'{cat["articulos"]} {n} {pieza}'
    frase = p.get("mensaje_web") or " ".join(x for x in [cuanto, vehiculo(p), anios(p)] if x)
    url = CONFIG["sitio_url"] + f'p/{p["id"]}/'
    wa = f"Hola, quería consultar por {frase} que vi en el catálogo (cód. {p['id']}).\n{url}"

    return {
        "id": p["id"],
        "tipo": p["tipo"],
        "titulo": titulo,
        "subtitulo": sub,
        "posicion": (p.get("posicion") or "").capitalize() or None,
        "cantidad": n,
        "cantidad_txt": p.get("cantidad_web") or etiqueta_cantidad(p, n),
        "estado": p.get("estado"),
        "segmentos": {"muy buena": 5, "buena": 4, "regular": 2}.get((p.get("estado") or "").lower()),
        "etiquetas": etiquetas,
        "sirve": sirve,
        "detalle": p.get("observaciones") or "",
        "precio_unidad": p.get("precio_unidad"),
        "precios_extra": precios_extra(p, n),
        "fotos": p.get("fotos", []),
        "nums": numeros(p),
        "texto": " ".join([titulo, sub, sirve, " ".join(etiquetas), " ".join(p.get("claves", []))]).lower(),
        "wa": wa,
        "url": url,
        "marca": p.get("marca"),
        "modelo": p.get("modelo"),
        "alta": p.get("fecha_ingreso"),
    }


# ---------- build ----------

def esc(s):
    return html.escape(str(s), quote=True)


def pagina_compartir(p, pub):
    """Página chica por producto: da la vista previa en WhatsApp y redirige al catálogo."""
    base = CONFIG["sitio_url"]
    destino = f"../../?p={p['id']}"
    if pub:
        titulo = f'{pub["titulo"]} {pub["subtitulo"]}'.strip()
        desc = " · ".join(x for x in [
            pub["cantidad_txt"],
            f'Estado: {pub["estado"]}' if pub["estado"] else "",
            f'{pesos(pub["precio_unidad"])}' if pub["precio_unidad"] else "Consultar precio",
        ] if x)
        img = base + pub["fotos"][0] if pub["fotos"] else base + "assets/og.jpg"
    else:
        titulo, desc, img = CONFIG["titulo"], "Este producto ya no está disponible. Mirá el resto del stock.", base + "assets/og.jpg"
    return f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<title>{esc(titulo)}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta property="og:type" content="product">
<meta property="og:title" content="{esc(titulo)}">
<meta property="og:description" content="{esc(desc)}">
<meta property="og:image" content="{esc(img)}">
<meta property="og:url" content="{esc(base + 'p/' + p['id'] + '/')}">
<meta name="robots" content="noindex">
<meta http-equiv="refresh" content="0; url={destino}">
<script>location.replace("{destino}")</script>
</head><body><a href="{destino}">Ver producto</a></body></html>
"""


def og_imagen():
    from PIL import Image, ImageDraw, ImageFont
    destino = ROOT / "assets" / "og.jpg"
    im = Image.new("RGB", (1200, 630), "#111111")
    d = ImageDraw.Draw(im)
    big = ImageFont.truetype(str(FONTS / "BarlowCondensed-Bold.ttf"), 112)
    mid = ImageFont.truetype(str(FONTS / "BarlowCondensed-Bold.ttf"), 58)
    small = ImageFont.truetype(str(FONTS / "Barlow-Medium.ttf"), 38)
    d.rectangle([0, 0, 1200, 18], fill="#FFC400")
    d.text((80, 100), "MDZAUTOPARTES\nCATÁLOGO", font=big, fill="white", spacing=4)
    d.text((80, 380), "PARAGOLPES, PARRILLAS Y ÓPTICAS USADAS", font=mid, fill="#FFC400")
    linea = CONFIG["ubicacion"] + (" · Hacemos envíos" if CONFIG.get("envios") else "")
    d.text((80, 490), linea, font=small, fill="#BDBDBD")
    im.save(destino, quality=82, optimize=True, progressive=True)


def build(silencioso=False):
    stock = cargar()
    productos = [publico(p) for p in stock["productos"] if visible(p)]
    productos.sort(key=lambda x: (x["alta"] or ""), reverse=True)
    for x in productos:
        x.pop("alta")
    salida = {
        "actualizado": ahora(),
        "sitio": {k: CONFIG[k] for k in ["titulo", "bajada", "whatsapp", "whatsapp_visible", "ubicacion", "envios", "sitio_url"]},
        "categorias": [{k: c[k] for k in ["id", "nombre", "singular", "plural", "filtro_posicion"]} for c in CONFIG["categorias"]],
        "productos": productos,
    }
    texto = json.dumps(salida, ensure_ascii=False, indent=1)
    for campo in PRIVADOS:  # control de seguridad
        if f'"{campo}"' in texto:
            sys.exit(f"ERROR: el campo privado '{campo}' iba a publicarse. Cancelado.")
    (ROOT / "data").mkdir(exist_ok=True)
    (ROOT / "data" / "productos.json").write_text(texto + "\n")

    carpeta = ROOT / "p"
    shutil.rmtree(carpeta, ignore_errors=True)
    pubs = {x["id"]: x for x in productos}
    for p in stock["productos"]:
        if p.get("prueba") and not visible(p):
            continue
        d = carpeta / p["id"]
        d.mkdir(parents=True, exist_ok=True)
        (d / "index.html").write_text(pagina_compartir(p, pubs.get(p["id"])))
    if not (ROOT / "assets" / "og.jpg").exists():
        og_imagen()
    if not silencioso:
        print(f"Build OK: {len(productos)} productos visibles.")


# ---------- fotos ----------

def procesar_fotos(pid, rutas, reemplazar=False):
    from PIL import Image, ImageOps
    stock = cargar()
    p = buscar(stock, pid)
    FOTOS.mkdir(exist_ok=True)
    if reemplazar:
        for f in p.get("fotos", []):
            (ROOT / f).unlink(missing_ok=True)
        p["fotos"] = []
    p.setdefault("fotos", [])
    usados = {int(m.group(1)) for f in p["fotos"] if (m := re.search(r"-(\d+)\.jpg$", f))}
    n = max(usados, default=0)
    for ruta in rutas:
        n += 1
        im = Image.open(ruta)
        im = ImageOps.exif_transpose(im).convert("RGB")  # endereza y descarta EXIF (GPS, etc.)
        im.thumbnail((1280, 1280), Image.LANCZOS)
        nombre = f"fotos/{p['id'].lower()}-{n}.jpg"
        im.save(ROOT / nombre, "JPEG", quality=72, optimize=True, progressive=True)
        p["fotos"].append(nombre)
        print(f"{nombre}  {(ROOT / nombre).stat().st_size // 1024} KB")
    guardar(stock)


# ---------- movimientos ----------

def opciones(args):
    pos, op = [], {}
    i = 0
    while i < len(args):
        a = args[i]
        if a.startswith("--"):
            clave = a[2:]
            if i + 1 < len(args) and not args[i + 1].startswith("--"):
                op[clave] = args[i + 1]
                i += 2
            else:
                op[clave] = True
                i += 1
        else:
            pos.append(a)
            i += 1
    return pos, op


def num(s):
    return float(str(s).replace("$", "").replace(".", "").replace(",", "."))


def cmd_nuevo(ruta):
    datos = json.loads(Path(ruta).read_text())
    stock = cargar()
    cat = categoria(datos["tipo"])
    sig = stock.setdefault("siguiente", {})
    n = sig.get(cat["prefijo"], 1)
    sig[cat["prefijo"]] = n + 1
    datos["id"] = f'{cat["prefijo"]}{n:03d}'
    datos.setdefault("activo", True)
    datos.setdefault("fecha_ingreso", hoy().isoformat())
    datos.setdefault("fotos", [])
    stock["productos"].append(datos)
    guardar(stock)
    registrar("ingreso", datos, int(datos["cantidad"]),
              int(datos["cantidad"]) * datos["costo"] if datos.get("costo") else None,
              f'Alta. {pesos(datos.get("precio_unidad"))} c/u' + (f', par {pesos(datos["precio_par"])}' if datos.get("precio_par") else ""))
    print(datos["id"])


def cmd_sumar(pid, n, op):
    stock = cargar()
    p = buscar(stock, pid)
    if "costo" in op:
        p["costo"] = num(op["costo"])
    p["cantidad"] = int(p["cantidad"]) + int(n)
    p["activo"] = True
    guardar(stock)
    registrar("ingreso", p, int(n), int(n) * p["costo"] if p.get("costo") else None,
              f'Se suman al lote. Ahora hay {p["cantidad"]}')
    print(f'{p["id"]}: ahora {p["cantidad"]}')


def cmd_venta(pid, n, op):
    stock = cargar()
    p = buscar(stock, pid)
    n = int(n)
    if n > int(p["cantidad"]):
        sys.exit(f'Solo hay {p["cantidad"]} de {p["id"]}')
    if "monto" in op:
        monto = num(op["monto"])
    elif p.get("precio_par") and n == 2:
        monto = p["precio_par"]
    else:
        monto = n * (p.get("precio_unidad") or 0)
    p["cantidad"] = int(p["cantidad"]) - n
    guardar(stock)
    registrar("venta", p, n, monto, op.get("nota", "") + (" — queda en 0, se oculta" if p["cantidad"] == 0 else ""))
    print(f'{p["id"]}: vendidas {n} por {pesos(monto)}. Quedan {p["cantidad"]}.')


def cmd_precio(pid, op):
    stock = cargar()
    p = buscar(stock, pid)
    cambios = []
    if "unidad" in op:
        cambios.append(f'unidad {pesos(p.get("precio_unidad"))} → {pesos(num(op["unidad"]))}')
        p["precio_unidad"] = num(op["unidad"])
    if "par" in op:
        if op["par"] in ("no", "0", "ninguno"):
            cambios.append("se quita precio por par")
            p.pop("precio_par", None)
        else:
            cambios.append(f'par {pesos(p.get("precio_par"))} → {pesos(num(op["par"]))}')
            p["precio_par"] = num(op["par"])
    guardar(stock)
    registrar("precio", p, None, None, "; ".join(cambios))
    print(f'{p["id"]}: ' + "; ".join(cambios))


def cmd_activo(pid, activo, op):
    stock = cargar()
    p = buscar(stock, pid)
    p["activo"] = activo
    guardar(stock)
    registrar("baja" if not activo else "alta", p, None, None, op.get("nota", ""))
    print(f'{p["id"]}: {"oculto" if not activo else "visible de nuevo"}')


def cmd_editar(pid, pares):
    stock = cargar()
    p = buscar(stock, pid)
    cambios = []
    for par in pares:
        k, v = par.split("=", 1)
        try:
            v = json.loads(v)
        except json.JSONDecodeError:
            pass
        if v is None:
            p.pop(k, None)
        else:
            p[k] = v
        cambios.append(f"{k}={v}")
    guardar(stock)
    registrar("edicion", p, None, None, "; ".join(cambios))
    print(f'{p["id"]}: ' + "; ".join(cambios))


def cmd_ver(texto, todos):
    stock = cargar()
    for p in stock["productos"]:
        if not todos and not visible(p):
            continue
        if texto and texto.lower() not in json.dumps(p, ensure_ascii=False).lower():
            continue
        estado = "" if visible(p) else "  [OCULTO]"
        print(f'{p["id"]}  {nombre_interno(p):50s} cant {p["cantidad"]}  {pesos(p.get("precio_unidad"))}'
              f'{"  par " + pesos(p["precio_par"]) if p.get("precio_par") else ""}'
              f'  {p.get("estado") or "-"}{estado}  fotos {len(p.get("fotos", []))}{"  PRUEBA" if p.get("prueba") else ""}')


def cmd_resumen(mes):
    stock = cargar()
    vis = [p for p in stock["productos"] if visible(p)]
    print(f"RESUMEN — {hoy().strftime('%d/%m/%Y')}\n")
    for cat in CONFIG["categorias"]:
        ps = [p for p in vis if p["tipo"] == cat["id"]]
        if not ps:
            continue
        unidades = sum(int(p["cantidad"]) for p in ps)
        valor = sum(int(p["cantidad"]) * (p.get("precio_unidad") or 0) for p in ps)
        print(f'{cat["nombre"].upper()}: {len(ps)} productos, {unidades} unidades, valor {pesos(valor)}')
        for p in ps:
            print(f'   {p["id"]}  {nombre_interno(p)} — {p["cantidad"]} u. x {pesos(p.get("precio_unidad"))}')
    total = sum(int(p["cantidad"]) * (p.get("precio_unidad") or 0) for p in vis)
    costo = sum(int(p["cantidad"]) * (p.get("costo") or 0) for p in vis)
    print(f"\nVALOR TOTAL DEL STOCK (a precio de venta): {pesos(total)}")
    if costo:
        print(f"Costo cargado del stock: {pesos(costo)}")
    mes = mes or hoy().strftime("%Y-%m")
    ventas = [h for h in historial() if h["tipo"] == "venta" and h["fecha"].startswith(mes)]
    print(f"\nVENTAS DE {mes}: {len(ventas)} operaciones, {sum(h['cantidad'] or 0 for h in ventas)} unidades, "
          f"{pesos(sum(h['monto'] or 0 for h in ventas))}")
    for h in ventas:
        print(f'   {h["fecha"][:10]}  {h["id"]} {h["producto"]} — {h["cantidad"]} u. {pesos(h["monto"])}')


PLANILLA = ROOT / "PLANILLA STOCK.xlsx"


def planilla(destino=None):
    """Planilla Excel siempre actualizada: stock con las mismas características que la web, movimientos y resumen."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    stock = cargar()
    hist = historial()
    wb = Workbook()

    def hoja(ws, cols, filas, dinero=()):
        ws.append(cols)
        for f in filas:
            ws.append(f)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="111111")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(42, max(10, largo + 2))
            if cols[i - 1] in dinero:
                for c in col[1:]:
                    c.number_format = '"$"#,##0'

    # --- Stock ---
    vendidas, ingresadas, ultima_venta = {}, {}, {}
    for m in hist:
        if m["tipo"] == "venta":
            vendidas[m["id"]] = vendidas.get(m["id"], 0) + (m["cantidad"] or 0)
            ultima_venta[m["id"]] = m["fecha"][:10]
        elif m["tipo"] == "ingreso":
            ingresadas[m["id"]] = ingresadas.get(m["id"], 0) + (m["cantidad"] or 0)
    cols = ["Código", "Nombre", "Tipo", "En la web", "Producto (web)", "Años / versión", "Marca auto", "Modelo",
            "Posición", "Lado", "Original / Alt.", "Estado", "Color", "N.º de pieza", "Sirve para", "Etiquetas (web)",
            "Observaciones (web)", "Ingresadas", "Vendidas", "Stock actual",
            "Precio ingreso c/u", "Precio venta c/u", "Precio par", "Ganancia c/u",
            "Valor stock (venta)", "Valor stock (costo)", "Fecha ingreso", "Última venta", "Link web",
            "Notas privadas", "Prueba"]
    filas = []
    for p in stock["productos"]:
        pub = publico(p)
        n = int(p["cantidad"])
        cat = categoria(p["tipo"])
        lote = ingresadas.get(p["id"], n + vendidas.get(p["id"], 0))
        filas.append([
            p["id"], nombre_planilla(p, lote), cat["singular"].capitalize(), "Sí" if visible(p) else "No",
            pub["titulo"], pub["subtitulo"], p.get("marca"), p.get("modelo"), p.get("posicion"), p.get("lado"),
            {True: "Original", False: "Alternativa"}.get(p.get("original")), p.get("estado"), p.get("color"),
            p.get("numero_pieza"), pub["sirve"] or None, ", ".join(pub["etiquetas"]) or None, p.get("observaciones"),
            lote, vendidas.get(p["id"], 0), n,
            p.get("costo"), p.get("precio_unidad"), p.get("precio_par"),
            (p["precio_unidad"] - p["costo"]) if p.get("precio_unidad") and p.get("costo") else None,
            n * (p.get("precio_unidad") or 0), n * (p.get("costo") or 0) if p.get("costo") else None,
            p.get("fecha_ingreso"), ultima_venta.get(p["id"]), pub["url"], p.get("notas_privadas"),
            "PRUEBA" if p.get("prueba") else "",
        ])
    ws = wb.active
    ws.title = "Stock"
    hoja(ws, cols, filas, dinero={"Precio ingreso c/u", "Precio venta c/u", "Precio par",
                                  "Ganancia c/u", "Valor stock (venta)", "Valor stock (costo)"})
    gris = PatternFill("solid", fgColor="E6E6E6")
    for fila in ws.iter_rows(min_row=2):
        if fila[3].value == "No":
            for c in fila:
                c.fill = gris

    # --- Movimientos ---
    nombres = {p["id"]: p for p in stock["productos"]}
    filas = []
    for m in hist:
        p = nombres.get(m["id"])
        cant = m.get("cantidad")
        unit = m["monto"] / cant if m.get("monto") is not None and cant else None
        ganancia = None
        if m["tipo"] == "venta" and m.get("monto") is not None and m.get("costo_unit"):
            ganancia = m["monto"] - cant * m["costo_unit"]
        filas.append([
            m["fecha"], {"ingreso": "INGRESO", "venta": "SALIDA (venta)", "baja": "BAJA", "alta": "ALTA",
                         "precio": "Cambio de precio", "edicion": "Corrección"}.get(m["tipo"], m["tipo"]),
            m["id"], nombre_planilla(p, cant or int(p["cantidad"])) if p else m["producto"], cant, unit, m.get("monto"),
            m.get("costo_unit"), ganancia, m.get("detalle"), "PRUEBA" if m.get("prueba") else "",
        ])
    hoja(wb.create_sheet("Movimientos"),
         ["Fecha", "Movimiento", "Código", "Nombre", "Cantidad", "Precio c/u (ingreso: costo · venta: cobrado)", "Monto total",
          "Costo c/u", "Ganancia", "Detalle", "Prueba"],
         filas, dinero={"Precio c/u (ingreso: costo · venta: cobrado)", "Monto total", "Costo c/u", "Ganancia"})

    # --- Resumen (sin productos de prueba) ---
    reales = [p for p in stock["productos"] if visible(p) and not p.get("prueba")]
    mes = hoy().strftime("%Y-%m")
    ventas_mes = [m for m in hist if m["tipo"] == "venta" and m["fecha"].startswith(mes) and not m.get("prueba")]
    filas = [
        ["Actualizado", ahora()],
        ["Productos publicados", len(reales)],
        ["Unidades en stock", sum(int(p["cantidad"]) for p in reales)],
        ["Valor del stock (precio de venta)", sum(int(p["cantidad"]) * (p.get("precio_unidad") or 0) for p in reales)],
        ["Valor del stock (precio de ingreso)", sum(int(p["cantidad"]) * (p.get("costo") or 0) for p in reales)],
        [f"Ventas de {mes} (unidades)", sum(m["cantidad"] or 0 for m in ventas_mes)],
        [f"Ventas de {mes} (monto)", sum(m["monto"] or 0 for m in ventas_mes)],
        [f"Ganancia de {mes}", sum((m["monto"] or 0) - (m["cantidad"] or 0) * (m.get("costo_unit") or 0)
                                    for m in ventas_mes if m.get("costo_unit"))],
        ["Nota", "Los totales no incluyen productos de prueba."],
    ]
    r = wb.create_sheet("Resumen", 0)
    hoja(r, ["Dato", "Valor"], filas)
    r.column_dimensions["A"].width = 38
    r.column_dimensions["B"].width = 28
    for fila in r.iter_rows(min_row=2):
        if "Valor" in fila[0].value or "monto" in fila[0].value or "Ganancia" in fila[0].value:
            fila[1].number_format = '"$"#,##0'
    r.freeze_panes = None

    destino = destino or PLANILLA
    try:
        wb.save(destino)
    except PermissionError:
        alt = destino.with_name(destino.stem + " (nueva).xlsx")
        wb.save(alt)
        print(f"AVISO: la planilla estaba abierta; se guardó como {alt.name}")
    return destino


def cmd_excel():
    (ROOT / "exportes").mkdir(exist_ok=True)
    ruta = planilla(ROOT / "exportes" / f"stock-{dt.datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx")
    planilla()
    print(ruta)


def cmd_publicar(mensaje):
    build()
    def git(*a, check=True):
        return subprocess.run(["git", *a], cwd=ROOT, check=check, capture_output=True, text=True)
    git("add", "-A")
    if not git("status", "--porcelain").stdout.strip():
        print("Sin cambios para publicar.")
        return
    git("commit", "-m", mensaje)
    r = git("push", check=False)
    if r.returncode:
        sys.exit("ERROR al publicar:\n" + r.stderr)
    print("Publicado:", mensaje)


def cmd_limpiar_pruebas():
    stock = cargar()
    quedan = []
    for p in stock["productos"]:
        if p.get("prueba"):
            for f in p.get("fotos", []):
                (ROOT / f).unlink(missing_ok=True)
            print("Borrado", p["id"], nombre_interno(p))
        else:
            quedan.append(p)
    stock["productos"] = quedan
    if not quedan:
        stock["siguiente"] = {}
    guardar(stock)
    HIST.write_text(json.dumps([h for h in historial() if not h.get("prueba")], ensure_ascii=False, indent=2) + "\n")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return
    cmd, (pos, op) = sys.argv[1], opciones(sys.argv[2:])
    if cmd == "ver":
        cmd_ver(pos[0] if pos else "", op.get("todos"))
    elif cmd == "nuevo":
        cmd_nuevo(pos[0])
    elif cmd == "sumar":
        cmd_sumar(pos[0], pos[1], op)
    elif cmd == "fotos":
        procesar_fotos(pos[0], pos[1:], op.get("reemplazar"))
    elif cmd == "venta":
        cmd_venta(pos[0], pos[1], op)
    elif cmd == "precio":
        cmd_precio(pos[0], op)
    elif cmd == "baja":
        cmd_activo(pos[0], False, op)
    elif cmd == "alta":
        cmd_activo(pos[0], True, op)
    elif cmd == "editar":
        cmd_editar(pos[0], pos[1:])
    elif cmd == "build":
        build()
    elif cmd == "resumen":
        cmd_resumen(op.get("mes"))
    elif cmd == "planilla":
        print(planilla())
    elif cmd == "excel":
        cmd_excel()
    elif cmd == "publicar":
        cmd_publicar(pos[0] if pos else "Actualización del catálogo")
    elif cmd == "limpiar-pruebas":
        cmd_limpiar_pruebas()
    else:
        print(__doc__)


if __name__ == "__main__":
    main()
